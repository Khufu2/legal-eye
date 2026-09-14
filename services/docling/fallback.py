"""Fallback extraction and low-memory content safety for Legal Eye processing.

Docling remains the preferred parser. The remote scanner remains the preferred
content-safety service. This module provides a deterministic low-memory fallback
so a scanner outage cannot strand otherwise parseable documents in the queue.
"""

from __future__ import annotations

import json
import zipfile
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

from pipeline import PipelineError


class _TextHTMLParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        text = data.strip()
        if text:
            self.parts.append(text)


def _split_text(value: str, max_chars: int = 2600) -> list[str]:
    paragraphs = [part.strip() for part in value.replace("\r\n", "\n").split("\n\n") if part.strip()]
    if not paragraphs:
        paragraphs = [part.strip() for part in value.split("\n") if part.strip()]
    chunks: list[str] = []
    current = ""
    for paragraph in paragraphs:
        if len(paragraph) > max_chars:
            if current:
                chunks.append(current)
                current = ""
            for start in range(0, len(paragraph), max_chars):
                piece = paragraph[start : start + max_chars].strip()
                if piece:
                    chunks.append(piece)
            continue
        candidate = f"{current}\n\n{paragraph}".strip() if current else paragraph
        if len(candidate) <= max_chars:
            current = candidate
        else:
            if current:
                chunks.append(current)
            current = paragraph
    if current:
        chunks.append(current)
    return chunks


def _scan_stream_for_tokens(path: Path, tokens: tuple[bytes, ...]) -> bytes | None:
    lowered = tuple(token.lower() for token in tokens)
    carry = b""
    longest = max((len(token) for token in lowered), default=1)
    with path.open("rb") as source:
        while block := source.read(1024 * 1024):
            window = (carry + block).lower()
            for token in lowered:
                if token in window:
                    return token
            carry = window[-(longest - 1):] if longest > 1 else b""
    return None


def _scan_ooxml(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            if not infos or len(infos) > 20_000:
                raise PipelineError("unsafe_archive", "Office document archive structure is invalid", False)
            total_uncompressed = 0
            blocked_suffixes = (
                "/vbaproject.bin", ".exe", ".dll", ".com", ".scr", ".ps1", ".vbs", ".vbe", ".js", ".jse", ".bat", ".cmd",
            )
            blocked_markers = ("/embeddings/", "/activex/", "/oleobject")
            for info in infos:
                normalized = info.filename.replace("\\", "/").lower()
                parts = [part for part in normalized.split("/") if part]
                if normalized.startswith("/") or ".." in parts:
                    raise PipelineError("unsafe_archive_path", "Office document contains an unsafe archive path", False)
                if info.flag_bits & 0x1:
                    raise PipelineError("encrypted_archive", "Encrypted Office documents are not accepted for automated processing", False)
                total_uncompressed += max(info.file_size, 0)
                if total_uncompressed > 250 * 1024 * 1024:
                    raise PipelineError("archive_too_large", "Expanded Office document exceeds the safety limit", False)
                if info.compress_size > 0 and info.file_size > 10 * 1024 * 1024 and info.file_size / info.compress_size > 250:
                    raise PipelineError("suspicious_compression", "Office document exceeds the compression-ratio safety limit", False)
                if normalized.endswith(blocked_suffixes) or any(marker in f"/{normalized}" for marker in blocked_markers):
                    raise PipelineError("active_content_blocked", "Office document contains active or embedded executable content", False)
    except PipelineError:
        raise
    except (OSError, zipfile.BadZipFile) as error:
        raise PipelineError("invalid_office_document", "Office document container is invalid", False) from error


def local_malware_scan(path: Path) -> dict[str, Any]:
    """Low-memory fallback gate for accepted Legal Eye document formats.

    This is deliberately described as content-safety scanning, not signature AV.
    It blocks active payloads and parser-abuse patterns while the remote scanner is
    unavailable, and records that provenance for later audit.
    """
    suffix = path.suffix.lower()
    size = path.stat().st_size
    if size <= 0 or size > 100 * 1024 * 1024:
        raise PipelineError("unsafe_file_size", "Document size is outside the content-safety policy", False)

    if suffix == ".pdf":
        with path.open("rb") as source:
            if source.read(5) != b"%PDF-":
                raise PipelineError("invalid_pdf_signature", "PDF signature is invalid", False)
        token = _scan_stream_for_tokens(path, (b"/JavaScript", b"/Launch", b"/EmbeddedFile", b"/RichMedia"))
        if token:
            raise PipelineError("active_content_blocked", f"PDF active content is blocked ({token.decode('ascii', 'ignore')})", False)
    elif suffix in {".docx", ".xlsx"}:
        _scan_ooxml(path)
    elif suffix in {".txt", ".html", ".htm"}:
        token = _scan_stream_for_tokens(path, (b"<script", b"javascript:", b"data:text/html"))
        if token:
            raise PipelineError("active_content_blocked", "Text document contains active web content", False)
    else:
        raise PipelineError("unsupported_safety_format", f"Content-safety fallback does not accept {suffix or 'this format'}", False)

    return {
        "scanner": "static-content-safety",
        "signature": "active-content-archive-policy-v1",
        "fallback": True,
        "signature_antivirus": False,
    }


def _text_nodes_from_pdf(path: Path) -> list[dict[str, Any]]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    nodes: list[dict[str, Any]] = []
    order = 0
    for page_index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        for part in _split_text(text):
            nodes.append({
                "text": part,
                "label": "paragraph",
                "self_ref": f"#/texts/{order}",
                "_source": "lightweight",
                "prov": [{"page_no": page_index, "confidence": 0.82}],
            })
            order += 1
    return nodes


def _text_nodes_from_docx(path: Path) -> list[dict[str, Any]]:
    from docx import Document

    document = Document(str(path))
    raw = "\n\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
    return [
        {
            "text": part,
            "label": "paragraph",
            "self_ref": f"#/texts/{index}",
            "_source": "lightweight",
            "prov": [{"confidence": 0.9}],
        }
        for index, part in enumerate(_split_text(raw))
    ]


def _text_nodes_from_xlsx(path: Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    nodes: list[dict[str, Any]] = []
    order = 0
    try:
        for sheet in workbook.worksheets:
            lines: list[str] = []
            for row in sheet.iter_rows(values_only=True):
                values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                if values:
                    lines.append("\t".join(values))
            for part in _split_text("\n".join(lines)):
                nodes.append({
                    "text": part,
                    "label": "table",
                    "self_ref": f"#/texts/{order}",
                    "_source": "lightweight",
                    "prov": [{"confidence": 0.88}],
                    "metadata": {"sheet": sheet.title},
                })
                order += 1
    finally:
        workbook.close()
    return nodes


def _text_nodes_from_plain(path: Path, content_type: str) -> list[dict[str, Any]]:
    text = path.read_text(encoding="utf-8", errors="replace")
    if "html" in content_type:
        parser = _TextHTMLParser()
        parser.feed(text)
        text = "\n\n".join(parser.parts)
    return [
        {
            "text": part,
            "label": "paragraph",
            "self_ref": f"#/texts/{index}",
            "_source": "lightweight",
            "prov": [{"confidence": 0.95}],
        }
        for index, part in enumerate(_split_text(text))
    ]


def lightweight_convert(path: Path, file_name: str, content_type: str, source_hash: str) -> tuple[dict[str, Any], bytes]:
    suffix = Path(file_name).suffix.lower()
    normalized = content_type.split(";", 1)[0].strip().lower()
    try:
        if suffix == ".pdf" or normalized == "application/pdf":
            texts = _text_nodes_from_pdf(path)
        elif suffix == ".docx" or normalized == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            texts = _text_nodes_from_docx(path)
        elif suffix == ".xlsx" or normalized == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
            texts = _text_nodes_from_xlsx(path)
        elif normalized in {"text/plain", "text/html", "application/xhtml+xml"} or suffix in {".txt", ".html", ".htm"}:
            texts = _text_nodes_from_plain(path, normalized)
        else:
            raise PipelineError("lightweight_format_unsupported", f"Lightweight extractor does not support {normalized or suffix}", False)
    except PipelineError:
        raise
    except Exception as error:
        raise PipelineError("lightweight_extract_failed", f"Lightweight extraction failed: {type(error).__name__}") from error

    if not texts:
        raise PipelineError("no_searchable_text", "No searchable text could be extracted from this document", False)

    payload = {
        "schema": "legal-eye.lightweight-conversion.v1",
        "engine": {"name": "lightweight", "version": "1.0.0"},
        "source": {"file_name": file_name, "content_type": content_type, "sha256": source_hash},
        "status": "success",
        "document": {"texts": texts, "tables": []},
        "provenance": {
            "confidence": {"text": 0.82},
            "canonical_format": "lightweight-json",
            "fallback": True,
        },
    }
    encoded = json.dumps(payload, separators=(",", ":"), default=str).encode("utf-8")
    return payload, encoded
